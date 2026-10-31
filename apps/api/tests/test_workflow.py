"""The whole product path, end to end, against a real database: upload, extract, evidence, draft,
verify, approve, export. Each step also has its negative case."""

import hashlib
import io
import json

import docx
import pytest
from conftest import SAMPLES

TENDER = (SAMPLES / "sample_tender.txt").read_bytes()
EVIDENCE = json.loads((SAMPLES / "sample_evidence.json").read_text())


def setup_tender(client, acct, upload=True):
    h = acct["headers"]
    tender = client.post("/v1/tenders", json={"title": "Riverside Council: Digital Services"}, headers=h).json()
    summary = None
    if upload:
        r = client.post(
            f"/v1/tenders/{tender['id']}/documents", headers=h, files={"file": ("tender.txt", TENDER, "text/plain")}
        )
        assert r.status_code == 201, r.text
        summary = r.json()
    return tender, summary


def add_evidence(client, acct, approved=True):
    ids = []
    for item in EVIDENCE:
        r = client.post("/v1/evidence", headers=acct["headers"], json={**item, "approved": approved})
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])
    return ids


def test_upload_extracts_verified_requirements(client, accounts):
    acct = accounts()
    tender, summary = setup_tender(client, acct)
    assert summary["kind"] == "txt" and summary["rejected_unverifiable"] == 0
    assert summary["questions"] == 2
    reqs = client.get(f"/v1/tenders/{tender['id']}/requirements", headers=acct["headers"]).json()
    types = {r["type"] for r in reqs}
    assert {"deadline", "scored_question", "mandatory_document", "declaration", "pass_fail"} <= types
    deadline = next(r for r in reqs if r["type"] == "deadline")
    assert deadline["due_at"].startswith("2026-11-14T12:00:00")  # 12:00 Europe/London (GMT in November)
    q1 = next(r for r in reqs if r["word_limit"] == 750)
    assert q1["weighting"] == 20.0
    for r in reqs:  # every requirement is verbatim text from the tender, with its source
        assert r["text_verbatim"] in TENDER.decode() and r["file_name"] == "tender.txt"


def test_questions_carry_the_full_prompt(client, accounts):
    acct = accounts()
    tender, _ = setup_tender(client, acct)
    qs = client.get(f"/v1/tenders/{tender['id']}/questions", headers=acct["headers"]).json()
    assert len(qs) == 2
    assert any("Describe your approach to mobilisation" in q["prompt"] and q["word_limit"] == 750 for q in qs)


def test_full_journey_to_a_gated_export(client, accounts):
    acct = accounts()
    h = acct["headers"]
    tender, _ = setup_tender(client, acct)
    add_evidence(client, acct)
    qs = client.get(f"/v1/tenders/{tender['id']}/questions", headers=h).json()
    drafts = [client.post(f"/v1/questions/{q['id']}/drafts", headers=h).json() for q in qs]
    assert all(d["claims"] and all(c["status"] == "supported" for c in d["claims"]) for d in drafts)
    assert any(d["evidence_used"] for d in drafts)
    # nothing is approved yet: the export is blocked
    blocked = client.get(f"/v1/tenders/{tender['id']}/export", headers=h)
    assert blocked.status_code == 409
    assert any("not approved" in p["problem"] for p in blocked.json()["detail"]["problems"])
    for d in drafts:
        assert client.post(f"/v1/drafts/{d['id']}/approve", headers=h).status_code == 200
    ready = client.get(f"/v1/tenders/{tender['id']}/readiness", headers=h).json()
    assert ready["ready"] is True
    ok = client.get(f"/v1/tenders/{tender['id']}/export", headers=h)
    assert ok.status_code == 200
    assert ok.headers["x-export-sha256"] == hashlib.sha256(ok.content).hexdigest()
    document = docx.Document(io.BytesIO(ok.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "AI Use & Verification Record" in text
    assert "[EV:" not in text  # evidence tags are stripped from the submission text
    record = document.tables[0]
    assert acct["email"] in {c.text for row in record.rows for c in row.cells}


def test_unsupported_claim_blocks_export_until_a_person_acknowledges_it(client, accounts):
    acct = accounts()
    h = acct["headers"]
    tender, _ = setup_tender(client, acct)
    add_evidence(client, acct)
    qs = client.get(f"/v1/tenders/{tender['id']}/questions", headers=h).json()
    drafts = [client.post(f"/v1/questions/{q['id']}/drafts", headers=h).json() for q in qs]
    first = drafts[0]
    edited = client.patch(
        f"/v1/drafts/{first['id']}", headers=h, json={"text": first["text"] + " We have delivered 40 projects."}
    ).json()
    assert edited["version"] == 2
    flagged = [c for c in edited["claims"] if c["status"] == "unsupported"]
    assert len(flagged) == 1 and "40 projects" in flagged[0]["sentence_span"]
    for d in [edited, *drafts[1:]]:
        client.post(f"/v1/drafts/{d['id']}/approve", headers=h)
    gate = client.get(f"/v1/tenders/{tender['id']}/export", headers=h)
    assert gate.status_code == 409 and "unsupported claim" in json.dumps(gate.json())
    assert client.post(f"/v1/claims/{flagged[0]['id']}/acknowledge", headers=h).status_code == 200
    assert client.get(f"/v1/tenders/{tender['id']}/export", headers=h).status_code == 200


def test_unapproved_evidence_is_never_used_or_trusted(client, accounts):
    acct = accounts()
    h = acct["headers"]
    tender, _ = setup_tender(client, acct)
    add_evidence(client, acct, approved=False)
    qs = client.get(f"/v1/tenders/{tender['id']}/questions", headers=h).json()
    d = client.post(f"/v1/questions/{qs[0]['id']}/drafts", headers=h).json()
    assert d["evidence_used"] == [] and "No approved evidence" in d["note"]


def test_provenance_tracks_origin_and_the_chain_detects_tampering(client, accounts, admin):
    acct = accounts()
    h = acct["headers"]
    tender, _ = setup_tender(client, acct)
    add_evidence(client, acct)
    qid = client.get(f"/v1/tenders/{tender['id']}/questions", headers=h).json()[0]["id"]
    d = client.post(f"/v1/questions/{qid}/drafts", headers=h).json()
    d2 = client.patch(f"/v1/drafts/{d['id']}", headers=h, json={"text": d["text"] + " Added by a person."}).json()
    prov = client.get(f"/v1/drafts/{d2['id']}/provenance", headers=h).json()
    origins = {s["origin"] for s in prov["spans"]}
    assert {"ai", "evidence", "human"} <= origins
    assert 0 < prov["ai_retained_ratio"] < 1
    assert prov["chain"] == {"ok": True, "first_bad_revision": None, "revisions": 2}
    # Prove privileged mutation is blocked with the actual append-only trigger enabled.
    import copy
    import uuid

    import psycopg
    from app.provenance.hashchain import Revision, verify_chain

    thread = uuid.UUID(d2["content"]["thread_id"]) if "content" in d2 else None
    if thread is None:
        content = admin.execute("SELECT content FROM drafts WHERE id=%s", (uuid.UUID(d2["id"]),)).fetchone()[0]
        thread = uuid.UUID(content["thread_id"])
    rows = admin.execute(
        "SELECT seq,prev_hash,hash,content_snapshot,actor_id FROM revisions WHERE draft_id=%s ORDER BY seq",
        (thread,),
    ).fetchall()
    captured = [Revision(copy.deepcopy(r[3]), {"actor": str(r[4]), "seq": r[0]}, r[1], r[2]) for r in rows]
    assert verify_chain(captured) == (True, None)
    with pytest.raises(psycopg.errors.RaiseException):
        admin.execute(
            "UPDATE revisions SET content_snapshot=jsonb_set(content_snapshot,'{text}','\"rewritten\"') "
            "WHERE org_id=%s AND seq=1",
            (acct["org_id"],),
        )
    # Verify a tampered immutable copy with the real verifier; never rewrite database evidence.
    changed = copy.deepcopy(captured)
    changed[0].content["text"] = "rewritten"
    assert verify_chain(changed) == (False, 0)
    intact = client.get(f"/v1/drafts/{d2['id']}/provenance", headers=h).json()["chain"]
    assert intact == {"ok": True, "first_bad_revision": None, "revisions": 2}


def test_stale_version_cannot_be_edited_or_approved(client, accounts):
    acct = accounts()
    h = acct["headers"]
    tender, _ = setup_tender(client, acct)
    qid = client.get(f"/v1/tenders/{tender['id']}/questions", headers=h).json()[0]["id"]
    v1 = client.post(f"/v1/questions/{qid}/drafts", headers=h).json()
    client.patch(f"/v1/drafts/{v1['id']}", headers=h, json={"text": "Second version."})
    assert client.patch(f"/v1/drafts/{v1['id']}", headers=h, json={"text": "x"}).status_code == 409
    assert client.post(f"/v1/drafts/{v1['id']}/approve", headers=h).status_code == 409


def test_word_limit_is_enforced_by_code(client, accounts):
    acct = accounts()
    h = acct["headers"]
    tender, _ = setup_tender(client, acct)
    qs = client.get(f"/v1/tenders/{tender['id']}/questions", headers=h).json()
    q2 = next(q for q in qs if q["word_limit"] == 500)
    d = client.post(f"/v1/questions/{q2['id']}/drafts", headers=h).json()
    long = client.patch(f"/v1/drafts/{d['id']}", headers=h, json={"text": "word " * 600}).json()
    assert long["over_limit"] is True and long["word_count"] == 600
    client.post(f"/v1/drafts/{long['id']}/approve", headers=h)
    problems = client.get(f"/v1/tenders/{tender['id']}/readiness", headers=h).json()["problems"]
    assert any("500-word limit" in p["problem"] for p in problems)


def test_audit_log_is_chained_and_verifiable(client, accounts):
    acct = accounts()
    h = acct["headers"]
    setup_tender(client, acct)
    log = client.get("/v1/audit", headers=h).json()
    assert log["chain_ok"] is True and log["events_total"] >= 3
    actions = [e["action"] for e in log["recent"]]
    assert "document.uploaded" in actions and "tender.created" in actions


def test_pdf_and_docx_uploads_are_parsed(client, accounts):
    import pymupdf

    acct = accounts()
    h = acct["headers"]
    tender = client.post("/v1/tenders", json={"title": "Formats"}, headers=h).json()
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 100), "Tenderers must provide a copy of their insurance certificate.", fontsize=11)
    r = client.post(
        f"/v1/tenders/{tender['id']}/documents", headers=h, files={"file": ("x.pdf", pdf.tobytes(), "application/pdf")}
    )
    assert r.status_code == 201 and r.json()["kind"] == "pdf" and r.json()["requirements"] == 1
    d = docx.Document()
    d.add_heading("Instructions", level=1)
    d.add_paragraph("Tenderers shall declare any conflicts of interest in writing.")
    buf = io.BytesIO()
    d.save(buf)
    r = client.post(
        f"/v1/tenders/{tender['id']}/documents",
        headers=h,
        files={"file": ("y.docx", buf.getvalue(), "application/octet-stream")},
    )
    assert r.status_code == 201 and r.json()["kind"] == "docx" and r.json()["requirements"] == 1
    reqs = client.get(f"/v1/tenders/{tender['id']}/requirements", headers=h).json()
    assert any(q["page"] == 1 and q["file_name"] == "x.pdf" for q in reqs)  # the source page is kept


@pytest.mark.parametrize(
    "name,payload",
    [
        ("fake.pdf", b"\x89PNG\r\n\x1a\n" + b"\x00" * 64),  # an image renamed to .pdf
        ("run.exe", b"MZ\x90\x00\x03" + b"\x00" * 64),  # an executable
        ("data.bin", bytes(range(256))),  # binary noise
        ("bad.docx", b"PK\x03\x04" + b"not really a zip"),  # a broken Office file
    ],
)
def test_disguised_or_unsupported_files_are_rejected(client, accounts, name, payload):
    acct = accounts()
    tender, _ = setup_tender(client, acct, upload=False)
    r = client.post(f"/v1/tenders/{tender['id']}/documents", headers=acct["headers"], files={"file": (name, payload)})
    assert r.status_code == 415


def test_oversized_and_empty_uploads_are_rejected(client, accounts, monkeypatch):
    from app.core.config import settings

    acct = accounts()
    tender, _ = setup_tender(client, acct, upload=False)
    url = f"/v1/tenders/{tender['id']}/documents"
    assert client.post(url, headers=acct["headers"], files={"file": ("e.txt", b"")}).status_code == 400
    monkeypatch.setattr(settings, "max_upload_mb", 0)
    assert client.post(url, headers=acct["headers"], files={"file": ("a.txt", b"hello")}).status_code == 413
