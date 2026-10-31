"""Safe document parsing: sniff the real file type, then extract text blocks with their location."""

import importlib
import io
import re
import zipfile

BLOCK_MIN_CHARS = 3


def initialize_parsers() -> None:
    """Load required codecs on the startup thread before accepting uploads.

    First import can initialize native runtimes. Keep that cost off the first
    request, and fail startup if a declared document capability is unavailable.
    PDF parsing stays on its existing single-threaded path.
    """
    importlib.import_module("pymupdf")
    importlib.import_module("docx")


def sniff(data: bytes) -> str | None:
    """Return 'pdf', 'docx' or 'txt' from the file's content, never from its name."""
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
            if "[Content_Types].xml" in names and "word/document.xml" in names:
                return "docx"
        except zipfile.BadZipFile:
            return None
        return None
    if b"\x00" in data[:4096]:
        return None
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        return None
    return "txt"


def parse(data: bytes, kind: str) -> list[dict]:
    """Return blocks: {page, kind, text, section_path, bbox}."""
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    return _txt(data)


def _pdf(data: bytes) -> list[dict]:
    import pymupdf

    blocks: list[dict] = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        for number, page in enumerate(doc, start=1):
            for x0, y0, x1, y1, text, *_ in page.get_text("blocks"):
                text = re.sub(r"\s+", " ", text).strip()
                if len(text) >= BLOCK_MIN_CHARS:
                    blocks.append(
                        {
                            "page": number,
                            "kind": "paragraph",
                            "text": text,
                            "section_path": None,
                            "bbox": [round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1)],
                        }
                    )
    return blocks


def _docx(data: bytes) -> list[dict]:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks, path = [], []
    for para in document.paragraphs:
        text = para.text.strip()
        if len(text) < BLOCK_MIN_CHARS:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        if style.lower().startswith("heading"):
            level = int(re.sub(r"\D", "", style) or 1)
            path = path[: level - 1] + [text]
            blocks.append(
                {"page": None, "kind": "heading", "text": text, "section_path": " > ".join(path), "bbox": None}
            )
        else:
            blocks.append(
                {
                    "page": None,
                    "kind": "paragraph",
                    "text": text,
                    "section_path": " > ".join(path) or None,
                    "bbox": None,
                }
            )
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if len(text) >= BLOCK_MIN_CHARS:
                blocks.append({"page": None, "kind": "table_row", "text": text, "section_path": None, "bbox": None})
    return blocks


def _txt(data: bytes) -> list[dict]:
    blocks, path = [], []
    for chunk in re.split(r"\n\s*\n", data.decode("utf-8")):
        text = re.sub(r"[ \t]*\n[ \t]*", " ", chunk).strip()
        if len(text) < BLOCK_MIN_CHARS:
            continue
        if len(text) < 90 and not re.search(r"[.!?:]$", text):
            path = [text]
            blocks.append({"page": None, "kind": "heading", "text": text, "section_path": text, "bbox": None})
        else:
            blocks.append(
                {
                    "page": None,
                    "kind": "paragraph",
                    "text": text,
                    "section_path": " > ".join(path) or None,
                    "bbox": None,
                }
            )
    return blocks
